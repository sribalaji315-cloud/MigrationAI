"""PowerBI xlsx ingestion — faithful Python port of the reference
``ImportPowerBi.ParseItemConfigurations`` / ``ParseFeatureGroups`` plus the
``PowerBiImportUtil`` cell helpers and ``ExcelExtensions.BuildColumnIndex``.

Uses openpyxl in read-only mode for bounded memory. Header names are matched
case-insensitively with Unicode whitespace normalized to a regular space (row 1
is the header). Empty cells return ``""`` (never ``None``) to match the
reference ``GetCellString`` semantics.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .expansion_engine import RawFeatureGroup, RawItemConfig


def _norm_ws(value: str) -> str:
    return "".join(" " if ch.isspace() else ch for ch in value)


def _build_col_index(header_row: Tuple) -> Dict[str, int]:
    idx: Dict[str, int] = {}
    for i, cell in enumerate(header_row):
        if cell is None:
            continue
        key = _norm_ws(str(cell).strip())
        if key:
            kl = key.lower()
            if kl not in idx:  # first occurrence wins
                idx[kl] = i
    return idx


def _cell_str(row: Tuple, col_index: Dict[str, int], name: str) -> str:
    i = col_index.get(name.lower())
    if i is None or i >= len(row):
        return ""
    v = row[i]
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()


def _cell_int(row: Tuple, col_index: Dict[str, int], name: str) -> int:
    i = col_index.get(name.lower())
    if i is None or i >= len(row):
        return 0
    v = row[i]
    if v is None:
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip()
    try:
        return int(s)
    except ValueError:
        try:
            return int(float(s))
        except ValueError:
            return 0


def _cell_datetime(row: Tuple, col_index: Dict[str, int], name: str) -> Optional[datetime]:
    i = col_index.get(name.lower())
    if i is None or i >= len(row):
        return None
    v = row[i]
    if v is None:
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = str(v).strip()
    if not s:
        return None
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y %H:%M:%S"):
            try:
                return datetime.strptime(s, fmt)
            except ValueError:
                continue
        return None


def _first_cell_str(row: Tuple) -> str:
    if len(row) == 0 or row[0] is None:
        return ""
    return str(row[0]).strip()


def parse_item_configs(path: str) -> List[RawItemConfig]:
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return []
        col = _build_col_index(header)

        items: List[RawItemConfig] = []
        last_item: Optional[str] = None
        last_desc: Optional[str] = None
        last_pg: Optional[str] = None

        for row in rows:
            if _first_cell_str(row).lower().startswith("applied filters"):
                break

            item_value = _cell_str(row, col, "Item")
            if item_value:
                last_item = item_value
                last_desc = _cell_str(row, col, "Description")
                last_pg = _cell_str(row, col, "Product Group")
            else:
                item_value = last_item or ""

            if not item_value:
                continue

            option = _cell_str(row, col, "Option")
            from_value = _cell_str(row, col, "From Value")
            to_value = _cell_str(row, col, "To Value")
            resolved_option = from_value if (not option and from_value) else option

            feature = _cell_str(row, col, "Feature")
            if not feature:
                continue

            desc = _cell_str(row, col, "Description")
            pg = _cell_str(row, col, "Product Group")

            items.append(
                RawItemConfig(
                    item=item_value,
                    description=desc or last_desc,
                    product_group=pg or last_pg,
                    feature=feature,
                    feature_description=_cell_str(row, col, "Feature Desc"),
                    option=resolved_option,
                    option_description=_cell_str(row, col, "Option Desc"),
                    sequence=_cell_int(row, col, "Seq."),
                    group=_cell_int(row, col, "Group"),
                    condition=_cell_str(row, col, "Condition"),
                    from_value=from_value,
                    to_value=to_value,
                    till=_cell_datetime(row, col, "Till"),
                )
            )
        return items
    finally:
        wb.close()


def parse_feature_groups(path: str) -> List[RawFeatureGroup]:
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return []
        col = _build_col_index(header)

        groups: List[RawFeatureGroup] = []
        for row in rows:
            if _first_cell_str(row).lower().startswith("applied filters"):
                break

            groups.append(
                RawFeatureGroup(
                    feature_group=_cell_str(row, col, "FeatureGroup"),
                    feature=_cell_str(row, col, "Feature"),
                    feature_description=_cell_str(row, col, "Feature Desc"),
                    option=_cell_str(row, col, "Option"),
                    option_description=_cell_str(row, col, "Option Desc"),
                    sequence=_cell_int(row, col, "Seq."),
                    group=_cell_int(row, col, "Group"),
                    condition=_cell_str(row, col, "Condition"),
                    till=_cell_datetime(row, col, "Till"),
                )
            )
        return groups
    finally:
        wb.close()
